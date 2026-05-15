#!/usr/bin/env python3
"""
Planilha de Acompanhamento Mensal de Rentabilidade
Modelo em branco para envio ao cliente: separa por classe de ativo (com
expansão dos ativos da classe), mês a mês, com valor bruto, valor líquido,
movimentação, rentabilidade em R$ e em %.
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────
# PARÂMETROS (ajuste conforme necessário)
# ─────────────────────────────────────────
ANO          = 2026
MES_INICIAL  = 1        # 1 = janeiro
QTD_MESES    = 12

# Classes de ativo: (nome, cor, qtde de linhas em branco para ativos)
# Para adicionar mais linhas a uma classe, basta aumentar o número.
# Para criar uma nova classe, adicione uma linha nesta lista.
CLASSES = [
    ("FUNDOS",                  "E6A817", 6),
    ("PREVIDÊNCIA INDIVIDUAL",  "6B7A1E", 6),
    ("CRÉDITO PRIVADO",         "2E9BD6", 6),
    ("OPERAÇÕES ESTRUTURADAS",  "BFA615", 4),
    ("CDB",                     "17B3C4", 6),
    ("LCI / LCA",               "2E8B57", 6),
    ("POUPANÇA",                "C0392B", 3),
]

# 5 colunas por mês
METRICAS = [
    "Valor Bruto",
    "Valor Líquido",
    "Movimentação\n(aporte + / resgate −)",
    "Rentabilidade\n(R$)",
    "Rentabilidade\n(%)",
]
N_MET = len(METRICAS)

MESES_ABREV = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun',
               'jul', 'ago', 'set', 'out', 'nov', 'dez']

# ─────────────────────────────────────────
# CORES / ESTILOS
# ─────────────────────────────────────────
NAVY     = '1F4E79'
AZUL_MED = '2E75B6'
CINZA_H  = 'F0F2F5'
ZEBRA    = 'F7F9FC'
BRANCO   = 'FFFFFF'

FMT_BRL  = 'R$ #,##0.00;[Red]-R$ #,##0.00'
FMT_PCT  = '0.00%;[Red]-0.00%'


def cor(hex_str):
    return PatternFill("solid", fgColor=hex_str)


def borda(claro=True):
    c = 'D9DEE5' if claro else 'AAB2BD'
    lado = Side(style='thin', color=c)
    return Border(left=lado, right=lado, top=lado, bottom=lado)


def cref(col, row):
    return f"{get_column_letter(col)}{row}"


# ─────────────────────────────────────────
# GERAÇÃO
# ─────────────────────────────────────────
def criar_planilha():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Acompanhamento Mensal"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.outlinePr.summaryBelow = False  # classe acima dos ativos

    # rótulos dos meses
    labels = []
    for i in range(QTD_MESES):
        idx = (MES_INICIAL - 1 + i) % 12
        ano = ANO + (MES_INICIAL - 1 + i) // 12
        labels.append(f"{MESES_ABREV[idx]}/{ano % 100:02d}")

    COL_LABEL = 1
    COL_SALDO = 2
    COL_M1    = 3
    TOTAL_COL = COL_M1 + QTD_MESES * N_MET - 1

    def mcol(m, off):           # m: 1..QTD_MESES ; off: 0..4
        return COL_M1 + (m - 1) * N_MET + off

    def prior_bruto_col(m):     # coluna do bruto do mês anterior
        return COL_SALDO if m == 1 else mcol(m - 1, 0)

    # ── larguras ──────────────────────────────────────────────────────────
    ws.column_dimensions[get_column_letter(COL_LABEL)].width = 36
    ws.column_dimensions[get_column_letter(COL_SALDO)].width = 15
    larguras = [13, 13, 14, 13, 11]
    for m in range(1, QTD_MESES + 1):
        for off in range(N_MET):
            ws.column_dimensions[get_column_letter(mcol(m, off))].width = larguras[off]

    # ── TÍTULO ────────────────────────────────────────────────────────────
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=TOTAL_COL)
    t = ws.cell(row=1, column=1, value="ACOMPANHAMENTO MENSAL DE RENTABILIDADE")
    t.fill = cor(NAVY)
    t.font = Font(bold=True, color=BRANCO, size=15)
    t.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 30

    def campo(row, rotulo):
        rc = ws.cell(row=row, column=1, value=rotulo)
        rc.font = Font(bold=True, size=10, color=NAVY)
        rc.alignment = Alignment(horizontal='left', vertical='center')
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        vc = ws.cell(row=row, column=2)
        vc.fill = cor('FFF7D6')
        vc.border = borda()
        vc.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.row_dimensions[row].height = 20

    campo(2, "Cliente:")
    campo(3, "Assessor / Instituição:")
    campo(4, "Período de referência:")

    # ── CABEÇALHO (linhas 6 e 7) ──────────────────────────────────────────
    HEAD1, HEAD2 = 6, 7

    ws.merge_cells(start_row=HEAD1, start_column=COL_LABEL,
                   end_row=HEAD2, end_column=COL_LABEL)
    h = ws.cell(row=HEAD1, column=COL_LABEL, value="CLASSE / ATIVO")
    h.fill = cor(NAVY)
    h.font = Font(bold=True, color=BRANCO, size=11)
    h.alignment = Alignment(horizontal='center', vertical='center')

    ws.merge_cells(start_row=HEAD1, start_column=COL_SALDO,
                   end_row=HEAD2, end_column=COL_SALDO)
    s = ws.cell(row=HEAD1, column=COL_SALDO, value="Saldo Inicial\n(Bruto)")
    s.fill = cor(NAVY)
    s.font = Font(bold=True, color=BRANCO, size=10)
    s.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for m in range(1, QTD_MESES + 1):
        c0 = mcol(m, 0)
        ws.merge_cells(start_row=HEAD1, start_column=c0,
                       end_row=HEAD1, end_column=c0 + N_MET - 1)
        mc = ws.cell(row=HEAD1, column=c0, value=labels[m - 1].upper())
        mc.fill = cor(NAVY if m % 2 else AZUL_MED)
        mc.font = Font(bold=True, color=BRANCO, size=11)
        mc.alignment = Alignment(horizontal='center', vertical='center')
        for off, met in enumerate(METRICAS):
            sc = ws.cell(row=HEAD2, column=c0 + off, value=met)
            sc.fill = cor(AZUL_MED)
            sc.font = Font(bold=True, color=BRANCO, size=8)
            sc.alignment = Alignment(horizontal='center', vertical='center',
                                     wrap_text=True)
            sc.border = borda()

    ws.row_dimensions[HEAD1].height = 22
    ws.row_dimensions[HEAD2].height = 42
    ws.freeze_panes = cref(COL_M1, HEAD2 + 1)

    # ── LINHAS DE DADOS ───────────────────────────────────────────────────
    PATRIM_ROW = HEAD2 + 1
    classe_rows = []
    classe_specs = []          # (header_row, primeiro_ativo, ultimo_ativo, cor)
    r = PATRIM_ROW + 1
    for nome, c, n in CLASSES:
        header = r
        primeiro = r + 1
        ultimo = r + n
        classe_rows.append(header)
        classe_specs.append((header, primeiro, ultimo, c, nome))
        r = ultimo + 1
    LAST_DATA = r - 1

    def estilo_celula(cell, num_fmt=None, bold=False, fill=None, fg=None):
        cell.border = borda()
        cell.alignment = Alignment(horizontal='right', vertical='center')
        cell.font = Font(size=10, bold=bold, color=fg or '000000')
        if fill:
            cell.fill = cor(fill)
        if num_fmt:
            cell.number_format = num_fmt

    # ---- PATRIMÔNIO TOTAL --------------------------------------------------
    ws.row_dimensions[PATRIM_ROW].height = 24
    pl = ws.cell(row=PATRIM_ROW, column=COL_LABEL, value="PATRIMÔNIO TOTAL")
    pl.fill = cor(NAVY)
    pl.font = Font(bold=True, color=BRANCO, size=11)
    pl.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    pl.border = borda()

    # saldo inicial total
    refs = ",".join(cref(COL_SALDO, h) for h in classe_rows)
    pc = ws.cell(row=PATRIM_ROW, column=COL_SALDO,
                 value=f'=IF(SUM({refs})=0,"",SUM({refs}))')
    estilo_celula(pc, FMT_BRL, bold=True, fill='D6E4F2', fg=NAVY)

    for m in range(1, QTD_MESES + 1):
        cb, cl, cm = mcol(m, 0), mcol(m, 1), mcol(m, 2)
        cr, cp = mcol(m, 3), mcol(m, 4)
        prior = prior_bruto_col(m)
        for col, fmt in ((cb, FMT_BRL), (cl, FMT_BRL), (cm, FMT_BRL),
                         (cr, FMT_BRL), (cp, FMT_PCT)):
            cell = ws.cell(row=PATRIM_ROW, column=col)
            if col in (cb, cl, cm, cr):
                rr = ",".join(cref(col, h) for h in classe_rows)
                cell.value = f'=IF(SUM({rr})=0,"",SUM({rr}))'
            else:  # rentabilidade %
                B = cref(cb, PATRIM_ROW)
                P = cref(prior, PATRIM_ROW)
                M = cref(cm, PATRIM_ROW)
                cell.value = (f'=IF(OR({B}="",{P}="",({P}+{M})=0),"",'
                              f'({B}-{P}-{M})/({P}+{M}))')
            estilo_celula(cell, fmt, bold=True, fill='D6E4F2', fg=NAVY)

    # ---- CLASSES + ATIVOS --------------------------------------------------
    for header, primeiro, ultimo, ccor, nome in classe_specs:
        # cabeçalho da classe
        ws.row_dimensions[header].height = 21
        hl = ws.cell(row=header, column=COL_LABEL, value=nome)
        hl.fill = cor(ccor)
        hl.font = Font(bold=True, color=BRANCO, size=10)
        hl.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        hl.border = borda()

        refs = ",".join(cref(COL_SALDO, rr) for rr in range(primeiro, ultimo + 1))
        hc = ws.cell(row=header, column=COL_SALDO,
                     value=f'=IF(SUM({refs})=0,"",SUM({refs}))')
        estilo_celula(hc, FMT_BRL, bold=True, fill=CINZA_H)

        for m in range(1, QTD_MESES + 1):
            cb, cl, cm = mcol(m, 0), mcol(m, 1), mcol(m, 2)
            cr, cp = mcol(m, 3), mcol(m, 4)
            prior = prior_bruto_col(m)
            for col, fmt in ((cb, FMT_BRL), (cl, FMT_BRL), (cm, FMT_BRL),
                             (cr, FMT_BRL), (cp, FMT_PCT)):
                cell = ws.cell(row=header, column=col)
                if col in (cb, cl, cr):
                    rng = f"{cref(col, primeiro)}:{cref(col, ultimo)}"
                    cell.value = f'=IF(SUM({rng})=0,"",SUM({rng}))'
                elif col == cm:
                    rng = f"{cref(col, primeiro)}:{cref(col, ultimo)}"
                    cell.value = f'=SUM({rng})'
                else:  # rent %
                    B = cref(cb, header)
                    P = cref(prior, header)
                    M = cref(cm, header)
                    cell.value = (f'=IF(OR({B}="",{P}="",({P}+{M})=0),"",'
                                  f'({B}-{P}-{M})/({P}+{M}))')
                estilo_celula(cell, fmt, bold=True, fill=CINZA_H)

        # linhas de ativos (em branco, agrupadas/recolhíveis)
        for idx, rr in enumerate(range(primeiro, ultimo + 1)):
            ws.row_dimensions[rr].outline_level = 1
            ws.row_dimensions[rr].height = 18
            zebra = ZEBRA if idx % 2 else BRANCO

            al = ws.cell(row=rr, column=COL_LABEL)
            al.fill = cor(zebra)
            al.border = borda()
            al.alignment = Alignment(horizontal='left', vertical='center', indent=3)
            al.font = Font(size=10, italic=True, color='8A93A0')

            sc = ws.cell(row=rr, column=COL_SALDO)
            estilo_celula(sc, FMT_BRL, fill=zebra)

            for m in range(1, QTD_MESES + 1):
                cb, cl, cm = mcol(m, 0), mcol(m, 1), mcol(m, 2)
                cr, cp = mcol(m, 3), mcol(m, 4)
                prior = prior_bruto_col(m)
                B, P, M = cref(cb, rr), cref(prior, rr), cref(cm, rr)
                # entradas
                for col in (cb, cl, cm):
                    estilo_celula(ws.cell(row=rr, column=col), FMT_BRL, fill=zebra)
                # rentabilidade R$
                cellr = ws.cell(row=rr, column=cr,
                                value=f'=IF({B}="","",{B}-{P}-{M})')
                estilo_celula(cellr, FMT_BRL, fill=zebra)
                # rentabilidade %
                cellp = ws.cell(row=rr, column=cp,
                                value=(f'=IF(OR({B}="",{P}="",({P}+{M})=0),"",'
                                       f'({B}-{P}-{M})/({P}+{M}))'))
                estilo_celula(cellp, FMT_PCT, fill=zebra)

    # ── FORMATAÇÃO CONDICIONAL (verde p/ ganho) ───────────────────────────
    verde = Font(color='1E7B34', bold=True)
    for m in range(1, QTD_MESES + 1):
        for off in (3, 4):
            col = get_column_letter(mcol(m, off))
            rng = f"{col}{PATRIM_ROW}:{col}{LAST_DATA}"
            ws.conditional_formatting.add(
                rng, CellIsRule(operator='greaterThan', formula=['0'], font=verde))

    # ══════════════════════════════════════════════════════════════════════
    # RESUMO + GRÁFICOS (mesma aba, abaixo da tabela)
    # ══════════════════════════════════════════════════════════════════════
    RES = LAST_DATA + 3
    ws.merge_cells(start_row=RES, start_column=1, end_row=RES, end_column=6)
    rt = ws.cell(row=RES, column=1, value="RESUMO – EVOLUÇÃO DO PATRIMÔNIO")
    rt.fill = cor(NAVY)
    rt.font = Font(bold=True, color=BRANCO, size=12)
    rt.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[RES].height = 24

    res_head = ["Mês", "Patrimônio Bruto", "Patrimônio Líquido",
                "Rentab. Mensal (R$)", "Rentab. Mensal (%)", "Rentab. Acum. (%)"]
    for col, txt in enumerate(res_head, 1):
        c = ws.cell(row=RES + 1, column=col, value=txt)
        c.fill = cor(AZUL_MED)
        c.font = Font(bold=True, color=BRANCO, size=9)
        c.alignment = Alignment(horizontal='center', vertical='center',
                                wrap_text=True)
        c.border = borda()
    ws.row_dimensions[RES + 1].height = 30

    primeira_res = RES + 2
    for i, m in enumerate(range(1, QTD_MESES + 1)):
        row = primeira_res + i
        zebra = ZEBRA if i % 2 else BRANCO
        b = cref(mcol(m, 0), PATRIM_ROW)
        l = cref(mcol(m, 1), PATRIM_ROW)
        rr = cref(mcol(m, 3), PATRIM_ROW)
        pp = cref(mcol(m, 4), PATRIM_ROW)

        valores = [
            labels[i],
            f'=IF({b}="","",{b})',
            f'=IF({l}="","",{l})',
            f'=IF({rr}="","",{rr})',
            f'=IF({pp}="","",{pp})',
        ]
        for col, val in enumerate(valores, 1):
            c = ws.cell(row=row, column=col, value=val)
            c.fill = cor(zebra)
            c.border = borda()
            c.font = Font(size=10)
            if col == 1:
                c.alignment = Alignment(horizontal='center', vertical='center')
            else:
                c.alignment = Alignment(horizontal='right', vertical='center')
                c.number_format = FMT_PCT if col == 5 else FMT_BRL
        # rentabilidade acumulada
        e_cur = cref(5, row)
        acc = ws.cell(row=row, column=6)
        if i == 0:
            acc.value = f'=IF({e_cur}="","",{e_cur})'
        else:
            e_prev = cref(6, row - 1)
            acc.value = (f'=IF({e_cur}="",{e_prev},IF({e_prev}="",{e_cur},'
                         f'(1+{e_prev})*(1+{e_cur})-1))')
        acc.fill = cor(zebra)
        acc.border = borda()
        acc.font = Font(size=10, bold=True)
        acc.alignment = Alignment(horizontal='right', vertical='center')
        acc.number_format = FMT_PCT

    ultima_res = primeira_res + QTD_MESES - 1
    for off in (3, 4, 5):
        col = get_column_letter(off + 1)
        ws.conditional_formatting.add(
            f"{col}{primeira_res}:{col}{ultima_res}",
            CellIsRule(operator='greaterThan', formula=['0'], font=verde))

    # ---- gráfico de linha: patrimônio bruto x líquido ---------------------
    linha = LineChart()
    linha.title = "Evolução do Patrimônio"
    linha.style = 12
    linha.y_axis.title = "R$"
    linha.y_axis.numFmt = 'R$ #,##0'
    linha.width, linha.height = 20, 10
    dados = Reference(ws, min_col=2, max_col=3,
                      min_row=RES + 1, max_row=ultima_res)
    cats = Reference(ws, min_col=1, min_row=primeira_res, max_row=ultima_res)
    linha.add_data(dados, titles_from_data=True)
    linha.set_categories(cats)
    ws.add_chart(linha, cref(8, RES))

    # ---- gráfico de barras: rentabilidade mensal --------------------------
    barra = BarChart()
    barra.title = "Rentabilidade Mensal (%)"
    barra.style = 10
    barra.y_axis.numFmt = '0.0%'
    barra.width, barra.height = 20, 10
    dados2 = Reference(ws, min_col=5, max_col=5,
                       min_row=RES + 1, max_row=ultima_res)
    barra.add_data(dados2, titles_from_data=True)
    barra.set_categories(cats)
    barra.legend = None
    ws.add_chart(barra, cref(8, RES + 21))

    # ══════════════════════════════════════════════════════════════════════
    # ABA INSTRUÇÕES
    # ══════════════════════════════════════════════════════════════════════
    wi = wb.create_sheet("Instruções")
    wi.sheet_view.showGridLines = False
    wi.column_dimensions['A'].width = 100

    wi.merge_cells('A1:A1')
    it = wi.cell(row=1, column=1, value="COMO USAR ESTA PLANILHA")
    it.fill = cor(NAVY)
    it.font = Font(bold=True, color=BRANCO, size=13)
    it.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    wi.row_dimensions[1].height = 28

    instrucoes = [
        "",
        "1. Preencha o nome do cliente, o assessor e o período no topo da aba "
        "'Acompanhamento Mensal'.",
        "",
        "2. Para cada classe, digite o nome dos ativos nas linhas em branco "
        "(coluna CLASSE / ATIVO).",
        "",
        "3. Em cada mês preencha apenas 3 colunas por ativo:",
        "     • Valor Bruto      – valor de mercado antes de impostos.",
        "     • Valor Líquido    – valor já descontado o IR no resgate.",
        "                          (em investimentos ISENTOS, repita o valor bruto)",
        "     • Movimentação     – aportes com sinal +  e  resgates com sinal −  "
        "(deixe vazio se não houve).",
        "",
        "4. As colunas Rentabilidade (R$) e Rentabilidade (%) são calculadas "
        "automaticamente:",
        "     Rentab. R$ = Valor Bruto do mês − Valor Bruto do mês anterior − "
        "Movimentação.",
        "     Assim, aportes e resgates NÃO são contados como rendimento.",
        "",
        "5. A coluna 'Saldo Inicial (Bruto)' é a base do 1º mês. Preencha-a com "
        "o valor antes do período acompanhado.",
        "",
        "6. As linhas de classe e a linha PATRIMÔNIO TOTAL somam tudo "
        "sozinhas — não digite nada nelas.",
        "",
        "7. Use os botões [-] / [+] à esquerda das linhas para recolher ou "
        "expandir os ativos de cada classe.",
        "",
        "8. O bloco RESUMO e os gráficos no fim da aba se atualizam sozinhos e "
        "podem ser usados no envio ao cliente.",
        "",
        "Observação: a diferença entre Valor Bruto e Valor Líquido mostra ao "
        "cliente quanto o IR impacta o resgate. Em produtos isentos (LCI, LCA, "
        "CRI, CRA, debêntures incentivadas, poupança) os dois valores são iguais.",
    ]
    for i, txt in enumerate(instrucoes, 2):
        c = wi.cell(row=i, column=1, value=txt)
        c.alignment = Alignment(horizontal='left', vertical='center',
                                wrap_text=True)
        negrito = txt[:2].strip().rstrip('.').isdigit()
        c.font = Font(size=11, bold=negrito,
                      color=NAVY if negrito else '333333')
        wi.row_dimensions[i].height = 19

    # ── salvar ────────────────────────────────────────────────────────────
    wb.active = ws
    path = "/home/user/Planilha/acompanhamento_mensal.xlsx"
    wb.save(path)
    print(f"Planilha gerada: {path}")


if __name__ == "__main__":
    criar_planilha()
