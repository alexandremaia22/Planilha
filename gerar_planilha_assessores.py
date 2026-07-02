#!/usr/bin/env python3
"""
Painel de Acompanhamento Diário — Equipe de Assessores de Investimentos
Compatível com Microsoft Excel e Google Sheets (sem macros).

Layout "largo": cada assessor ocupa 1 única linha por mês; os dias viram
blocos de colunas (um bloco de 9 colunas por dia). O Saldo Anterior é
puxado automaticamente do Saldo Pulso do dia anterior (ou do último dia do
mês anterior / Saldo Inicial, no caso do dia 1), então o assessor só
digita: Saldo Pulso, Captação Líquida, COE, Previdência, Crédito Privado,
Títulos Públicos e Observações.
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.styles.protection import Protection
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter

# ─────────────────────────────────────────────────────────────────────────
# PARÂMETROS GERAIS
# ─────────────────────────────────────────────────────────────────────────
NUM_ASSESSORES = 13
ANO_PADRAO = 2026
DIAS_MAX = 31

CONFIG_SHEET = "Configurações"
CONFIG_YEAR_ROW = 3
CONFIG_HEADER_ROW = 6
CONFIG_DATA_START = CONFIG_HEADER_ROW + 1                   # 7
CONFIG_DATA_END = CONFIG_DATA_START + NUM_ASSESSORES - 1    # 19
CONFIG_COL_NOME = 2       # B
CONFIG_COL_META = 3       # C
CONFIG_COL_SALDO_INI = 4  # D

MESES = [
    (1, "Janeiro"), (2, "Fevereiro"), (3, "Março"), (4, "Abril"),
    (5, "Maio"), (6, "Junho"), (7, "Julho"), (8, "Agosto"),
    (9, "Setembro"), (10, "Outubro"), (11, "Novembro"), (12, "Dezembro"),
]

# ── Paleta corporativa sóbria ───────────────────────────────────────────
AZUL_ESCURO    = "1F3864"
AZUL           = "2E5395"
AZUL_CLARO     = "DCE6F1"
CINZA_CLARO    = "F2F2F2"
CINZA_BLOCO    = "E7E6E6"
BRANCO         = "FFFFFF"
CINZA_TEXTO    = "595959"
VERDE          = "1E7145"
VERDE_CLARO    = "E2F0D9"
VERMELHO       = "C00000"
VERMELHO_CLARO = "FCE4E4"
OURO           = "FFD966"
PRATA          = "D9D9D9"
BRONZE         = "E4C9A6"
LARANJA        = "C55A11"
AMARELO_INPUT  = "FFFDE7"
CINZA_INVALIDO = "D9D9D9"

BRL_FMT = 'R$ #,##0.00;[RED]-R$ #,##0.00'
BRL_FMT_SIMPLE = 'R$ #,##0.00'

# ── Estrutura dos blocos diários (formato largo) ────────────────────────
# cada dia = 9 colunas: Saldo Pulso, Saldo Anterior, Captação Líquida,
# COE, Previdência, Crédito Privado, Títulos Públicos, Produção Total, Obs.
CAMPOS = ["Saldo Pulso", "Saldo Anterior", "Captação Líq.", "COE",
          "Previdência", "Créd. Privado", "Tít. Públicos", "Produção Tot.", "Obs."]
N_CAMPOS = len(CAMPOS)
OFF_SP, OFF_SA, OFF_CL, OFF_COE, OFF_PREV, OFF_CP, OFF_TP, OFF_PT, OFF_OBS = range(N_CAMPOS)
CAMPOS_EDITAVEIS = {OFF_SP, OFF_CL, OFF_COE, OFF_PREV, OFF_CP, OFF_TP, OFF_OBS}

COL_ASSESSOR = 1                      # coluna A
FIRST_DAY_COL = 2                     # coluna B = início do bloco do dia 1
SUMMARY_START = FIRST_DAY_COL + DIAS_MAX * N_CAMPOS   # 2 + 279 = 281

SUM_PRODTOTAL, SUM_CAPTLIQ, SUM_SALDOATUAL, SUM_MEDIA, SUM_META, SUM_STATUS, SUM_CHAVE = range(7)
SUMMARY_LABELS = ["Produção Total (Mês)", "Captação Líquida (Mês)", "Saldo Pulso Atual",
                   "Média Diária", "Meta Mensal", "Status Meta", "Chave"]

LAST_COL = SUMMARY_START + len(SUMMARY_LABELS) - 1  # última coluna usada (inclui chave oculta)

# Linhas fixas do layout de cada aba mensal
ROW_TITLE = 1
ROW_INFO = 2
ROW_KPI_TITLE = 4
ROW_KPI_LABEL = 5
ROW_KPI_VALUE = 6
ROW_PROD_TITLE = 8
ROW_PROD_LABEL = 9
ROW_PROD_VALUE = 10
ROW_RANK_TITLE = 12
ROW_RANK_HEADER = 13
ROW_RANK_START = 14
ROW_RANK_END = ROW_RANK_START + NUM_ASSESSORES - 1     # 26

ROW_CHART_1 = 28
ROW_CHART_2 = 46

ROW_TABLE_TITLE = 64
ROW_BLOCK_HEADER = 65        # "Dia 1 — 01/02 (Qui)" mesclado por bloco
ROW_FIELD_HEADER = 66        # cabeçalho dos 9 campos, repetido por bloco
ROW_DATA_START = 67
ROW_DATA_END = ROW_DATA_START + NUM_ASSESSORES - 1     # 79

FREEZE_CELL = f"B{ROW_DATA_START}"

# Tabela auxiliar oculta — evolução diária da equipe (para o gráfico de linha)
EVOL_COL_DIA = LAST_COL + 2
EVOL_COL_DATA = LAST_COL + 3
EVOL_COL_PROD = LAST_COL + 4
EVOL_COL_CAPT = LAST_COL + 5
EVOL_HEADER_ROW = 1
EVOL_START = 2
EVOL_END = EVOL_START + DIAS_MAX - 1  # 32

PROD_LBL_COL = LAST_COL + 7
PROD_VAL_COL = LAST_COL + 8
PROD_START = 2
PROD_END = 5  # 4 produtos


def col(idx):
    return get_column_letter(idx)


def block_col(d, offset):
    """Coluna (índice) do campo `offset` dentro do bloco do dia `d` (1..31)."""
    return FIRST_DAY_COL + (d - 1) * N_CAMPOS + offset


def summary_col(offset):
    return SUMMARY_START + offset


# ─────────────────────────────────────────────────────────────────────────
# HELPERS DE ESTILO
# ─────────────────────────────────────────────────────────────────────────
def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)


def thin_border(color="D9D9D9"):
    s = Side(style='thin', color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def box_border(color="BFBFBF", style='medium'):
    s = Side(style=style, color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def style_header_cell(cell, bg=AZUL_ESCURO, fg=BRANCO, size=10):
    cell.fill = fill(bg)
    cell.font = Font(bold=True, color=fg, size=size, name="Calibri")
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border("FFFFFF")


def style_title_cell(cell, text, bg=AZUL_ESCURO, fg=BRANCO, size=14):
    cell.value = text
    cell.fill = fill(bg)
    cell.font = Font(bold=True, color=fg, size=size, name="Calibri")
    cell.alignment = Alignment(horizontal='left', vertical='center', indent=1)


def style_kpi_label(cell, text):
    cell.value = text
    cell.fill = fill(AZUL)
    cell.font = Font(bold=True, color=BRANCO, size=9, name="Calibri")
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border("FFFFFF")


def style_kpi_value(cell):
    cell.fill = fill(AZUL_CLARO)
    cell.font = Font(bold=True, color=AZUL_ESCURO, size=12, name="Calibri")
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border(BRANCO)


def lock(cell, locked=True):
    cell.protection = Protection(locked=locked)


def protect_sheet(ws):
    ws.protection.sheet = True
    ws.protection.formatCells = False
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.autoFilter = False
    ws.protection.sort = False
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False


# ─────────────────────────────────────────────────────────────────────────
# ABA DE CONFIGURAÇÃO
# ─────────────────────────────────────────────────────────────────────────
def montar_config(wb):
    ws = wb.create_sheet(CONFIG_SHEET)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 24
    ws.column_dimensions['D'].width = 20
    ws.column_dimensions['E'].width = 3
    ws.column_dimensions['F'].width = 70

    ws.merge_cells('A1:D1')
    style_title_cell(ws['A1'], "⚙  CONFIGURAÇÕES GERAIS — PAINEL DE ACOMPANHAMENTO DE ASSESSORES", size=13)
    ws.row_dimensions[1].height = 30

    ws.merge_cells('A2:D2')
    info = ws['A2']
    info.value = ("Preencha o ano, o nome dos 13 assessores, a meta mensal e o saldo pulso inicial (opcional). "
                  "Todo o restante — dias do mês, totais, ranking, saldos e gráficos — é calculado automaticamente nas 12 abas mensais.")
    info.font = Font(italic=True, size=9, color=CINZA_TEXTO)
    info.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[2].height = 28

    lbl = ws.cell(row=CONFIG_YEAR_ROW, column=1, value="Ano de Referência")
    ws.merge_cells(start_row=CONFIG_YEAR_ROW, start_column=1, end_row=CONFIG_YEAR_ROW, end_column=1)
    lbl.fill = fill(AZUL_CLARO)
    lbl.font = Font(bold=True, size=11, color=AZUL_ESCURO)
    lbl.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    lbl.border = thin_border()

    val = ws.cell(row=CONFIG_YEAR_ROW, column=2, value=ANO_PADRAO)
    val.fill = fill(BRANCO)
    val.font = Font(bold=True, size=12, color=VERDE)
    val.alignment = Alignment(horizontal='center', vertical='center')
    val.border = box_border(AZUL, 'medium')
    val.number_format = '0'
    lock(val, False)
    ws.row_dimensions[CONFIG_YEAR_ROW].height = 24

    dv_ano = DataValidation(type="whole", operator="between", formula1=2000, formula2=2100,
                             allow_blank=False, showErrorMessage=True,
                             errorTitle="Ano inválido", error="Informe um ano entre 2000 e 2100.")
    ws.add_data_validation(dv_ano)
    dv_ano.add(val)

    for c, (txt, w) in enumerate([("Nº", 6), ("Nome do Assessor", 30), ("Meta Mensal de Produção (R$)", 24),
                                   ("Saldo Pulso Inicial (R$)", 20)], 1):
        cell = ws.cell(row=CONFIG_HEADER_ROW, column=c, value=txt)
        style_header_cell(cell, bg=AZUL_ESCURO)
    ws.row_dimensions[CONFIG_HEADER_ROW].height = 26

    dv_meta = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0,
                              allow_blank=True, showErrorMessage=True,
                              errorTitle="Valor inválido", error="Informe um valor numérico maior ou igual a zero.")
    ws.add_data_validation(dv_meta)
    dv_saldo = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0,
                               allow_blank=True, showErrorMessage=True,
                               errorTitle="Valor inválido", error="Informe um valor numérico maior ou igual a zero.")
    ws.add_data_validation(dv_saldo)

    for i in range(1, NUM_ASSESSORES + 1):
        row = CONFIG_DATA_START + i - 1
        bg = CINZA_CLARO if i % 2 == 0 else BRANCO
        ws.row_dimensions[row].height = 20

        n = ws.cell(row=row, column=1, value=i)
        n.alignment = Alignment(horizontal='center', vertical='center')
        n.font = Font(bold=True, size=10, color=AZUL_ESCURO)
        n.fill = fill(bg)
        n.border = thin_border()

        nome = ws.cell(row=row, column=CONFIG_COL_NOME, value=None)
        nome.fill = fill(BRANCO)
        nome.font = Font(size=10, bold=True)
        nome.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        nome.border = box_border(AZUL, 'thin')
        lock(nome, False)

        meta = ws.cell(row=row, column=CONFIG_COL_META, value=None)
        meta.fill = fill(BRANCO)
        meta.font = Font(size=10)
        meta.alignment = Alignment(horizontal='right', vertical='center')
        meta.number_format = BRL_FMT_SIMPLE
        meta.border = box_border(AZUL, 'thin')
        lock(meta, False)
        dv_meta.add(meta)

        saldo_ini = ws.cell(row=row, column=CONFIG_COL_SALDO_INI, value=None)
        saldo_ini.fill = fill(BRANCO)
        saldo_ini.font = Font(size=10)
        saldo_ini.alignment = Alignment(horizontal='right', vertical='center')
        saldo_ini.number_format = BRL_FMT_SIMPLE
        saldo_ini.border = box_border(AZUL, 'thin')
        lock(saldo_ini, False)
        dv_saldo.add(saldo_ini)

    nota_row = CONFIG_DATA_END + 2
    ws.merge_cells(f'A{nota_row}:D{nota_row}')
    n1 = ws.cell(row=nota_row, column=1,
                 value="🔒  As células em branco com borda azul são as únicas editáveis. O restante da planilha é protegido contra alterações acidentais.")
    n1.font = Font(italic=True, size=9, color=CINZA_TEXTO)
    n1.alignment = Alignment(horizontal='left', wrap_text=True)
    ws.row_dimensions[nota_row].height = 26

    nota_row2 = nota_row + 1
    ws.merge_cells(f'A{nota_row2}:D{nota_row2}')
    n2 = ws.cell(row=nota_row2, column=1,
                 value="ℹ  O Saldo Pulso Inicial é usado apenas como Saldo Anterior do dia 1 de Janeiro. Nos demais meses, o Saldo Anterior do dia 1 puxa automaticamente o último Saldo Pulso do mês anterior.")
    n2.font = Font(italic=True, size=9, color=CINZA_TEXTO)
    n2.alignment = Alignment(horizontal='left', wrap_text=True)
    ws.row_dimensions[nota_row2].height = 26

    protect_sheet(ws)
    return ws


# ─────────────────────────────────────────────────────────────────────────
# CONSTRUÇÃO DE UMA ABA MENSAL
# ─────────────────────────────────────────────────────────────────────────
def montar_mes(wb, mnum, mnome, mes_anterior):
    ws = wb.create_sheet(mnome)
    ws.sheet_view.showGridLines = False

    ws.column_dimensions['A'].width = 22
    campo_larguras = [13, 13, 14, 11, 12, 13, 13, 13, 20]
    for d in range(1, DIAS_MAX + 1):
        for off, w in enumerate(campo_larguras):
            ws.column_dimensions[col(block_col(d, off))].width = w

    for off, (label, w) in enumerate(zip(SUMMARY_LABELS, [16, 16, 16, 14, 14, 14, 10])):
        ws.column_dimensions[col(summary_col(off))].width = w
    ws.column_dimensions[col(SUM_CHAVE + SUMMARY_START)].hidden = True

    for c in [EVOL_COL_DIA, EVOL_COL_DATA, EVOL_COL_PROD, EVOL_COL_CAPT, PROD_LBL_COL, PROD_VAL_COL]:
        ws.column_dimensions[col(c)].width = 16
        ws.column_dimensions[col(c)].hidden = True

    ANO = f"{CONFIG_SHEET}!$B${CONFIG_YEAR_ROW}"
    DIAS_MES = "$Z$3"     # célula auxiliar: dias no mês
    LAST_LETTER = col(LAST_COL)

    # ══ TÍTULO ══
    ws.merge_cells(f'A{ROW_TITLE}:{LAST_LETTER}{ROW_TITLE}')
    t = ws[f'A{ROW_TITLE}']
    t.value = f'="ACOMPANHAMENTO DIÁRIO DE ASSESSORES — {mnome.upper()} / "&{ANO}'
    t.fill = fill(AZUL_ESCURO)
    t.font = Font(bold=True, color=BRANCO, size=15, name="Calibri")
    t.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[ROW_TITLE].height = 32

    ws.merge_cells(f'A{ROW_INFO}:{LAST_LETTER}{ROW_INFO}')
    info = ws[f'A{ROW_INFO}']
    info.value = (f'="Dias no mês: "&{DIAS_MES}&"   |   Cada assessor ocupa 1 linha. Role para a direita para ver os dias. '
                  f'Apenas as células destacadas em amarelo são editáveis."')
    info.font = Font(italic=True, size=9, color=CINZA_TEXTO)
    info.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[ROW_INFO].height = 18

    # helper: dias no mês
    ws['Y3'] = "Dias no mês"
    ws['Y3'].font = Font(size=8, italic=True)
    ws['Z3'] = f"=DAY(EOMONTH(DATE({ANO},{mnum},1),0))"
    ws['Z3'].font = Font(size=8, bold=True)

    # ══ KPI CARDS ══
    ws.merge_cells(f'A{ROW_KPI_TITLE}:L{ROW_KPI_TITLE}')
    style_title_cell(ws[f'A{ROW_KPI_TITLE}'], "📊  INDICADORES DO MÊS", bg=AZUL, size=11)
    ws.row_dimensions[ROW_KPI_TITLE].height = 22

    SPT = col(summary_col(SUM_PRODTOTAL))
    SCL = col(summary_col(SUM_CAPTLIQ))
    SSA = col(summary_col(SUM_SALDOATUAL))

    kpi_defs = [
        ("PRODUÇÃO TOTAL DO MÊS", f"=SUM(${SPT}${ROW_DATA_START}:${SPT}${ROW_DATA_END})", BRL_FMT_SIMPLE),
        ("CAPTAÇÃO LÍQUIDA DO MÊS", f"=SUM(${SCL}${ROW_DATA_START}:${SCL}${ROW_DATA_END})", BRL_FMT_SIMPLE),
        ("SALDO PULSO TOTAL (ATUAL)", f"=SUM(${SSA}${ROW_DATA_START}:${SSA}${ROW_DATA_END})", BRL_FMT_SIMPLE),
        ("MÉDIA DIÁRIA DA EQUIPE", f"=IFERROR(SUM(${SPT}${ROW_DATA_START}:${SPT}${ROW_DATA_END})/{DIAS_MES},0)", BRL_FMT_SIMPLE),
        ("MELHOR DIA DE PRODUÇÃO",
         f'=IFERROR("Dia "&TEXT(INDEX(${col(EVOL_COL_DATA)}${EVOL_START}:${col(EVOL_COL_DATA)}${EVOL_END},'
         f'MATCH(MAX(${col(EVOL_COL_PROD)}${EVOL_START}:${col(EVOL_COL_PROD)}${EVOL_END}),'
         f'${col(EVOL_COL_PROD)}${EVOL_START}:${col(EVOL_COL_PROD)}${EVOL_END},0)),"dd/mm")&"  -  "&'
         f'TEXT(MAX(${col(EVOL_COL_PROD)}${EVOL_START}:${col(EVOL_COL_PROD)}${EVOL_END}),"R$ #,##0"),"—")',
         None),
        ("MELHOR ASSESSOR DO MÊS",
         f'=IFERROR($B${ROW_RANK_START}&"  -  "&TEXT($C${ROW_RANK_START},"R$ #,##0"),"—")', None),
    ]
    span = 12 // len(kpi_defs)
    for i, (label, formula, fmt) in enumerate(kpi_defs):
        c0 = i * span + 1
        c1 = c0 + span - 1
        ws.merge_cells(start_row=ROW_KPI_LABEL, start_column=c0, end_row=ROW_KPI_LABEL, end_column=c1)
        style_kpi_label(ws.cell(row=ROW_KPI_LABEL, column=c0), label)
        ws.merge_cells(start_row=ROW_KPI_VALUE, start_column=c0, end_row=ROW_KPI_VALUE, end_column=c1)
        vcell = ws.cell(row=ROW_KPI_VALUE, column=c0)
        vcell.value = formula
        style_kpi_value(vcell)
        if fmt:
            vcell.number_format = fmt
        else:
            vcell.font = Font(bold=True, color=AZUL_ESCURO, size=10, name="Calibri")
    ws.row_dimensions[ROW_KPI_LABEL].height = 26
    ws.row_dimensions[ROW_KPI_VALUE].height = 30

    # ══ TOTAIS POR PRODUTO ══
    ws.merge_cells(f'A{ROW_PROD_TITLE}:L{ROW_PROD_TITLE}')
    style_title_cell(ws[f'A{ROW_PROD_TITLE}'], "💰  TOTAIS POR PRODUTO", bg=AZUL, size=11)
    ws.row_dimensions[ROW_PROD_TITLE].height = 22

    produtos = [("COE", OFF_COE), ("PREVIDÊNCIA", OFF_PREV), ("CRÉDITO PRIVADO", OFF_CP), ("TÍTULOS PÚBLICOS", OFF_TP)]
    span2 = 12 // len(produtos)
    for i, (nome, offset) in enumerate(produtos):
        c0 = i * span2 + 1
        c1 = c0 + span2 - 1
        ws.merge_cells(start_row=ROW_PROD_LABEL, start_column=c0, end_row=ROW_PROD_LABEL, end_column=c1)
        style_kpi_label(ws.cell(row=ROW_PROD_LABEL, column=c0), nome)
        ws.merge_cells(start_row=ROW_PROD_VALUE, start_column=c0, end_row=ROW_PROD_VALUE, end_column=c1)
        vcell = ws.cell(row=ROW_PROD_VALUE, column=c0)
        termos = "+".join(f"SUM(${col(block_col(d, offset))}${ROW_DATA_START}:${col(block_col(d, offset))}${ROW_DATA_END})"
                           for d in range(1, DIAS_MAX + 1))
        vcell.value = f"={termos}"
        style_kpi_value(vcell)
        vcell.number_format = BRL_FMT_SIMPLE

        ws.cell(row=PROD_START + i, column=PROD_LBL_COL, value=nome)
        pv = ws.cell(row=PROD_START + i, column=PROD_VAL_COL)
        pv.value = f"={termos}"
    ws.row_dimensions[ROW_PROD_LABEL].height = 20
    ws.row_dimensions[ROW_PROD_VALUE].height = 26

    # ══ RANKING DE ASSESSORES ══
    ws.merge_cells(f'A{ROW_RANK_TITLE}:L{ROW_RANK_TITLE}')
    style_title_cell(ws[f'A{ROW_RANK_TITLE}'], "🏆  RANKING DE ASSESSORES (TOP 5 EM DESTAQUE)", bg=AZUL, size=11)
    ws.row_dimensions[ROW_RANK_TITLE].height = 22

    rank_headers = ["Pos.", "Assessor", "Produção Total", "Captação Líquida",
                     "Saldo Pulso Atual", "Média Diária", "Meta Mensal", "Status"]
    for i, htxt in enumerate(rank_headers, 1):
        c = ws.cell(row=ROW_RANK_HEADER, column=i, value=htxt)
        style_header_cell(c, bg=AZUL_ESCURO)
    ws.row_dimensions[ROW_RANK_HEADER].height = 22

    SCH = col(summary_col(SUM_CHAVE))
    SMD = col(summary_col(SUM_MEDIA))
    SMT = col(summary_col(SUM_META))

    for p in range(1, NUM_ASSESSORES + 1):
        row = ROW_RANK_START + p - 1
        ws.row_dimensions[row].height = 18
        pos_c = ws.cell(row=row, column=1, value=p)
        pos_c.alignment = Alignment(horizontal='center', vertical='center')
        pos_c.font = Font(bold=True, size=10)
        pos_c.border = thin_border()

        idxf = f"MATCH(LARGE(${SCH}${ROW_DATA_START}:${SCH}${ROW_DATA_END},{p}),${SCH}${ROW_DATA_START}:${SCH}${ROW_DATA_END},0)"
        nome_c = ws.cell(row=row, column=2, value=f"=INDEX($A${ROW_DATA_START}:$A${ROW_DATA_END},{idxf})")
        nome_c.font = Font(bold=True, size=10)
        nome_c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        nome_c.border = thin_border()

        prod_c = ws.cell(row=row, column=3, value=f"=INDEX(${SPT}${ROW_DATA_START}:${SPT}${ROW_DATA_END},{idxf})")
        prod_c.number_format = BRL_FMT_SIMPLE
        prod_c.alignment = Alignment(horizontal='right', vertical='center')
        prod_c.border = thin_border()

        capt_c = ws.cell(row=row, column=4, value=f"=INDEX(${SCL}${ROW_DATA_START}:${SCL}${ROW_DATA_END},{idxf})")
        capt_c.number_format = BRL_FMT
        capt_c.alignment = Alignment(horizontal='right', vertical='center')
        capt_c.border = thin_border()

        saldo_c = ws.cell(row=row, column=5, value=f"=INDEX(${SSA}${ROW_DATA_START}:${SSA}${ROW_DATA_END},{idxf})")
        saldo_c.number_format = BRL_FMT_SIMPLE
        saldo_c.alignment = Alignment(horizontal='right', vertical='center')
        saldo_c.border = thin_border()

        media_c = ws.cell(row=row, column=6, value=f"=INDEX(${SMD}${ROW_DATA_START}:${SMD}${ROW_DATA_END},{idxf})")
        media_c.number_format = BRL_FMT_SIMPLE
        media_c.alignment = Alignment(horizontal='right', vertical='center')
        media_c.border = thin_border()

        meta_c = ws.cell(row=row, column=7, value=f"=INDEX(${SMT}${ROW_DATA_START}:${SMT}${ROW_DATA_END},{idxf})")
        meta_c.number_format = BRL_FMT_SIMPLE
        meta_c.alignment = Alignment(horizontal='right', vertical='center')
        meta_c.border = thin_border()

        status_c = ws.cell(row=row, column=8,
                            value=f'=IF(G{row}=0,"— Sem meta",IF(C{row}>=G{row},"✓ Meta atingida","Em andamento"))')
        status_c.alignment = Alignment(horizontal='center', vertical='center')
        status_c.font = Font(size=9, bold=True)
        status_c.border = thin_border()

        bg = CINZA_CLARO if p % 2 == 0 else BRANCO
        cells = (pos_c, nome_c, prod_c, capt_c, saldo_c, media_c, meta_c, status_c)
        for cc in cells:
            cc.fill = fill(bg)
        if p == 1:
            for cc in cells: cc.fill = fill(OURO)
        elif p == 2:
            for cc in cells: cc.fill = fill(PRATA)
        elif p == 3:
            for cc in cells: cc.fill = fill(BRONZE)
        elif p <= 5:
            for cc in cells: cc.fill = fill(VERDE_CLARO)

    ws.conditional_formatting.add(
        f"H{ROW_RANK_START}:H{ROW_RANK_END}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("Atingida",H{ROW_RANK_START}))'], fill=fill(VERDE_CLARO), font=Font(color=VERDE, bold=True))
    )

    # ── Tabela auxiliar de evolução diária (oculta) — fonte do gráfico de linha
    ws.cell(row=EVOL_HEADER_ROW, column=EVOL_COL_DIA, value="Dia")
    ws.cell(row=EVOL_HEADER_ROW, column=EVOL_COL_DATA, value="Data")
    ws.cell(row=EVOL_HEADER_ROW, column=EVOL_COL_PROD, value="Produção do Dia")
    ws.cell(row=EVOL_HEADER_ROW, column=EVOL_COL_CAPT, value="Captação do Dia")
    for d in range(1, DIAS_MAX + 1):
        row = EVOL_START + d - 1
        ws.cell(row=row, column=EVOL_COL_DIA, value=d).font = Font(size=8)
        zc = ws.cell(row=row, column=EVOL_COL_DATA, value=f'=IF({d}<={DIAS_MES},DATE({ANO},{mnum},{d}),"")')
        zc.number_format = 'dd/mm'
        zc.font = Font(size=8)
        pc = col(block_col(d, OFF_PT))
        cc_ = col(block_col(d, OFF_CL))
        prodc = ws.cell(row=row, column=EVOL_COL_PROD,
                         value=f"=SUM(${pc}${ROW_DATA_START}:${pc}${ROW_DATA_END})")
        prodc.font = Font(size=8)
        captc = ws.cell(row=row, column=EVOL_COL_CAPT,
                         value=f"=SUM(${cc_}${ROW_DATA_START}:${cc_}${ROW_DATA_END})")
        captc.font = Font(size=8)

    # ══ GRÁFICOS ══
    montar_graficos(ws)

    # ══ TABELA DE LANÇAMENTOS DIÁRIOS (FORMATO LARGO) ══
    ws.merge_cells(f'A{ROW_TABLE_TITLE}:{LAST_LETTER}{ROW_TABLE_TITLE}')
    style_title_cell(ws[f'A{ROW_TABLE_TITLE}'], "📋  LANÇAMENTOS DIÁRIOS POR ASSESSOR (role para a direita)", bg=AZUL_ESCURO, size=11)
    ws.row_dimensions[ROW_TABLE_TITLE].height = 22

    # cabeçalho fixo da coluna Assessor
    ws.merge_cells(start_row=ROW_BLOCK_HEADER, start_column=1, end_row=ROW_FIELD_HEADER, end_column=1)
    ac = ws.cell(row=ROW_BLOCK_HEADER, column=1, value="Assessor")
    style_header_cell(ac, bg=AZUL_ESCURO, size=11)
    ws.row_dimensions[ROW_BLOCK_HEADER].height = 20
    ws.row_dimensions[ROW_FIELD_HEADER].height = 26

    # cabeçalho de cada bloco de dia
    for d in range(1, DIAS_MAX + 1):
        c0 = block_col(d, 0)
        c1 = block_col(d, N_CAMPOS - 1)
        ws.merge_cells(start_row=ROW_BLOCK_HEADER, start_column=c0, end_row=ROW_BLOCK_HEADER, end_column=c1)
        hcell = ws.cell(row=ROW_BLOCK_HEADER, column=c0)
        hcell.value = (f'="Dia {d} — "&IF({d}<={DIAS_MES},TEXT(DATE({ANO},{mnum},{d}),"dd/mm")&'
                        f'" ("&CHOOSE(WEEKDAY(DATE({ANO},{mnum},{d})),"Dom","Seg","Ter","Qua","Qui","Sex","Sáb")&")","—")')
        style_header_cell(hcell, bg=AZUL, size=9)
        for off, label in enumerate(CAMPOS):
            fc = ws.cell(row=ROW_FIELD_HEADER, column=c0 + off, value=label)
            style_header_cell(fc, bg=AZUL_ESCURO, size=8)

    # cabeçalho das colunas de resumo mensal
    for off, label in enumerate(SUMMARY_LABELS):
        hc = ws.cell(row=ROW_BLOCK_HEADER, column=summary_col(off), value=label if off != SUM_CHAVE else "")
        style_header_cell(hc, bg=LARANJA, size=8)
        ws.merge_cells(start_row=ROW_BLOCK_HEADER, start_column=summary_col(off),
                        end_row=ROW_FIELD_HEADER, end_column=summary_col(off))

    # validações de dados (uma por coluna-tipo, aplicadas a todos os 31 blocos de uma vez)
    dv_saldo = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0,
                               allow_blank=True, showErrorMessage=True,
                               errorTitle="Valor inválido", error="Informe um valor numérico maior ou igual a zero.")
    dv_capt = DataValidation(type="decimal", operator="between", formula1=-999999999, formula2=999999999,
                              allow_blank=True, showErrorMessage=True,
                              errorTitle="Valor inválido", error="Informe um valor numérico (positivo ou negativo).")
    dv_prod = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1=0,
                              allow_blank=True, showErrorMessage=True,
                              errorTitle="Valor inválido", error="Informe um valor numérico maior ou igual a zero.")
    dv_obs = DataValidation(type="textLength", operator="lessThanOrEqual", formula1=300,
                             allow_blank=True, showErrorMessage=True,
                             errorTitle="Texto muito longo", error="Limite de 300 caracteres.")
    for dv in (dv_saldo, dv_capt, dv_prod, dv_obs):
        ws.add_data_validation(dv)
    for d in range(1, DIAS_MAX + 1):
        dv_saldo.add(f"{col(block_col(d, OFF_SP))}{ROW_DATA_START}:{col(block_col(d, OFF_SP))}{ROW_DATA_END}")
        dv_capt.add(f"{col(block_col(d, OFF_CL))}{ROW_DATA_START}:{col(block_col(d, OFF_CL))}{ROW_DATA_END}")
        for off in (OFF_COE, OFF_PREV, OFF_CP, OFF_TP):
            dv_prod.add(f"{col(block_col(d, off))}{ROW_DATA_START}:{col(block_col(d, off))}{ROW_DATA_END}")
        dv_obs.add(f"{col(block_col(d, OFF_OBS))}{ROW_DATA_START}:{col(block_col(d, OFF_OBS))}{ROW_DATA_END}")

    # linhas de dados — 1 linha por assessor
    for i in range(1, NUM_ASSESSORES + 1):
        row = ROW_DATA_START + i - 1
        cfg_row = CONFIG_DATA_START + i - 1
        bg = CINZA_CLARO if i % 2 == 0 else BRANCO
        ws.row_dimensions[row].height = 20

        nome_ref = f"{CONFIG_SHEET}!$B${cfg_row}"
        nome_c = ws.cell(row=row, column=1, value=f'=IF({nome_ref}="","Assessor {i}",{nome_ref})')
        nome_c.font = Font(bold=True, size=10)
        nome_c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        nome_c.fill = fill(AZUL_CLARO)
        nome_c.border = box_border(AZUL, 'thin')

        for d in range(1, DIAS_MAX + 1):
            c0 = block_col(d, 0)
            sp_c = ws.cell(row=row, column=c0 + OFF_SP)
            sp_c.number_format = BRL_FMT_SIMPLE
            lock(sp_c, False)
            sp_c.fill = fill(AMARELO_INPUT)

            if d == 1:
                if mes_anterior is None:
                    sa_formula = f"={CONFIG_SHEET}!$D${cfg_row}"
                else:
                    prev_saldo_col = col(summary_col(SUM_SALDOATUAL))
                    sa_formula = f"='{mes_anterior}'!${prev_saldo_col}{row}"
            else:
                prev_sp_col = col(block_col(d - 1, OFF_SP))
                sa_formula = f"={prev_sp_col}{row}"
            sa_c = ws.cell(row=row, column=c0 + OFF_SA, value=sa_formula)
            sa_c.number_format = BRL_FMT_SIMPLE
            sa_c.font = Font(size=10, color=CINZA_TEXTO)

            cl_c = ws.cell(row=row, column=c0 + OFF_CL)
            cl_c.number_format = BRL_FMT
            lock(cl_c, False)
            cl_c.fill = fill(AMARELO_INPUT)

            for off in (OFF_COE, OFF_PREV, OFF_CP, OFF_TP):
                pc = ws.cell(row=row, column=c0 + off)
                pc.number_format = BRL_FMT_SIMPLE
                lock(pc, False)
                pc.fill = fill(AMARELO_INPUT)

            coe_l = col(c0 + OFF_COE)
            prev_l = col(c0 + OFF_PREV)
            cp_l = col(c0 + OFF_CP)
            tp_l = col(c0 + OFF_TP)
            pt_c = ws.cell(row=row, column=c0 + OFF_PT,
                            value=f'=SUM({coe_l}{row}:{tp_l}{row})')
            pt_c.number_format = BRL_FMT_SIMPLE
            pt_c.font = Font(bold=True, size=10, color=AZUL_ESCURO)

            obs_c = ws.cell(row=row, column=c0 + OFF_OBS)
            lock(obs_c, False)
            obs_c.fill = fill(AMARELO_INPUT)
            obs_c.alignment = Alignment(horizontal='left', vertical='center', indent=1)

            for offx in range(N_CAMPOS):
                cell = ws.cell(row=row, column=c0 + offx)
                cell.border = thin_border()
                if offx not in (OFF_SP, OFF_CL, OFF_COE, OFF_PREV, OFF_CP, OFF_TP, OFF_OBS):
                    cell.fill = fill(bg)
                if offx != OFF_OBS:
                    cell.alignment = Alignment(horizontal='right', vertical='center')

        # ── colunas de resumo mensal por assessor ──
        termos_pt = "+".join(f"{col(block_col(d, OFF_PT))}{row}" for d in range(1, DIAS_MAX + 1))
        pt_tot = ws.cell(row=row, column=summary_col(SUM_PRODTOTAL), value=f"={termos_pt}")
        pt_tot.number_format = BRL_FMT_SIMPLE

        termos_cl = "+".join(f"{col(block_col(d, OFF_CL))}{row}" for d in range(1, DIAS_MAX + 1))
        cl_tot = ws.cell(row=row, column=summary_col(SUM_CAPTLIQ), value=f"={termos_cl}")
        cl_tot.number_format = BRL_FMT

        opcoes_sp = ", ".join(f"{col(block_col(d, OFF_SP))}{row}" for d in range(1, DIAS_MAX + 1))
        sa_tot = ws.cell(row=row, column=summary_col(SUM_SALDOATUAL), value=f"=CHOOSE({DIAS_MES},{opcoes_sp})")
        sa_tot.number_format = BRL_FMT_SIMPLE

        media_tot = ws.cell(row=row, column=summary_col(SUM_MEDIA),
                             value=f"=IFERROR({col(summary_col(SUM_PRODTOTAL))}{row}/{DIAS_MES},0)")
        media_tot.number_format = BRL_FMT_SIMPLE

        meta_tot = ws.cell(row=row, column=summary_col(SUM_META), value=f"={CONFIG_SHEET}!$C${cfg_row}")
        meta_tot.number_format = BRL_FMT_SIMPLE

        status_tot = ws.cell(row=row, column=summary_col(SUM_STATUS),
                              value=(f'=IF({col(summary_col(SUM_META))}{row}=0,"— Sem meta",'
                                     f'IF({col(summary_col(SUM_PRODTOTAL))}{row}>={col(summary_col(SUM_META))}{row},'
                                     f'"✓ Atingida","Em andamento"))'))
        status_tot.alignment = Alignment(horizontal='center', vertical='center')
        status_tot.font = Font(size=9, bold=True)

        chave = ws.cell(row=row, column=summary_col(SUM_CHAVE),
                         value=f"={col(summary_col(SUM_PRODTOTAL))}{row}+({NUM_ASSESSORES + 1 - i}*0.0000001)")

        for offc in range(len(SUMMARY_LABELS)):
            cc = ws.cell(row=row, column=summary_col(offc))
            cc.fill = fill("FCE9D8")
            cc.border = thin_border()
            if offc not in (SUM_STATUS,):
                cc.alignment = Alignment(horizontal='right', vertical='center')

    # ══ FORMATAÇÃO CONDICIONAL ══
    last_field_letter = col(block_col(DIAS_MAX, N_CAMPOS - 1))
    # esmaece blocos de dias que não existem no mês corrente
    ws.conditional_formatting.add(
        f"B{ROW_BLOCK_HEADER}:{last_field_letter}{ROW_DATA_END}",
        FormulaRule(formula=[f'{DIAS_MES}<INT((COLUMN()-{FIRST_DAY_COL})/{N_CAMPOS})+1'],
                    fill=fill(CINZA_INVALIDO), font=Font(color='9B9B9B', italic=True))
    )
    for d in range(1, DIAS_MAX + 1):
        cl_letter = col(block_col(d, OFF_CL))
        rng = f"{cl_letter}{ROW_DATA_START}:{cl_letter}{ROW_DATA_END}"
        ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThan', formula=['0'],
                                                        font=Font(color=VERDE, bold=True)))
        ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['0'],
                                                        font=Font(color=VERMELHO, bold=True)))
        pt_letter = col(block_col(d, OFF_PT))
        ws.conditional_formatting.add(f"{pt_letter}{ROW_DATA_START}:{pt_letter}{ROW_DATA_END}",
                                       ColorScaleRule(start_type='min', start_color='FFFFFF', end_type='max', end_color='63BE7B'))

    ws.conditional_formatting.add(
        f"{col(summary_col(SUM_STATUS))}{ROW_DATA_START}:{col(summary_col(SUM_STATUS))}{ROW_DATA_END}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("Atingida",{col(summary_col(SUM_STATUS))}{ROW_DATA_START}))'],
                    fill=fill(VERDE_CLARO), font=Font(color=VERDE, bold=True))
    )

    # ══ FILTRO, CONGELAMENTO, PROTEÇÃO ══
    ws.auto_filter.ref = f"A{ROW_FIELD_HEADER}:{col(summary_col(SUM_STATUS))}{ROW_DATA_END}"
    ws.freeze_panes = FREEZE_CELL

    ws.sheet_view.zoomScale = 90
    protect_sheet(ws)

    ws.page_setup.orientation = 'landscape'
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    return ws


def montar_graficos(ws):
    bar = BarChart()
    bar.type = "bar"
    bar.style = 10
    bar.title = "Ranking de Assessores — Produção Total"
    bar.x_axis.title = "R$"
    bar.width = 17
    bar.height = 15
    bar.legend = None
    data = Reference(ws, min_col=3, min_row=ROW_RANK_HEADER, max_row=ROW_RANK_END)
    cats = Reference(ws, min_col=2, min_row=ROW_RANK_START, max_row=ROW_RANK_END)
    bar.add_data(data, titles_from_data=True)
    bar.set_categories(cats)
    bar.dataLabels = DataLabelList()
    bar.dataLabels.showVal = False
    if bar.series:
        bar.series[0].graphicalProperties.solidFill = AZUL
    ws.add_chart(bar, f"A{ROW_CHART_1}")

    colc = BarChart()
    colc.type = "col"
    colc.style = 10
    colc.title = "Comparativo — Captação Líquida por Assessor"
    colc.y_axis.title = "R$"
    colc.width = 17
    colc.height = 15
    colc.legend = None
    data2 = Reference(ws, min_col=4, min_row=ROW_RANK_HEADER, max_row=ROW_RANK_END)
    cats2 = Reference(ws, min_col=2, min_row=ROW_RANK_START, max_row=ROW_RANK_END)
    colc.add_data(data2, titles_from_data=True)
    colc.set_categories(cats2)
    if colc.series:
        colc.series[0].graphicalProperties.solidFill = LARANJA
    ws.add_chart(colc, f"G{ROW_CHART_1}")

    line = LineChart()
    line.title = "Evolução Diária — Produção e Captação"
    line.style = 10
    line.y_axis.title = "R$"
    line.x_axis.title = "Dia do mês"
    line.width = 17
    line.height = 15
    dataP = Reference(ws, min_col=EVOL_COL_PROD, min_row=EVOL_HEADER_ROW, max_row=EVOL_END)
    dataC = Reference(ws, min_col=EVOL_COL_CAPT, min_row=EVOL_HEADER_ROW, max_row=EVOL_END)
    catsD = Reference(ws, min_col=EVOL_COL_DIA, min_row=EVOL_START, max_row=EVOL_END)
    line.add_data(dataP, titles_from_data=True)
    line.add_data(dataC, titles_from_data=True)
    line.set_categories(catsD)
    if len(line.series) >= 2:
        line.series[0].graphicalProperties.line.solidFill = VERDE
        line.series[0].smooth = False
        line.series[1].graphicalProperties.line.solidFill = LARANJA
        line.series[1].smooth = False
    ws.add_chart(line, f"A{ROW_CHART_2}")

    pie = PieChart()
    pie.title = "Participação por Produto"
    pie.style = 10
    pie.width = 17
    pie.height = 15
    dataPie = Reference(ws, min_col=PROD_VAL_COL, min_row=PROD_START, max_row=PROD_END)
    catsPie = Reference(ws, min_col=PROD_LBL_COL, min_row=PROD_START, max_row=PROD_END)
    pie.add_data(dataPie, titles_from_data=False)
    pie.set_categories(catsPie)
    pie.dataLabels = DataLabelList()
    pie.dataLabels.showPercent = True
    ws.add_chart(pie, f"G{ROW_CHART_2}")


# ─────────────────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────────────────
def gerar():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    montar_config(wb)
    mes_anterior = None
    for mnum, mnome in MESES:
        montar_mes(wb, mnum, mnome, mes_anterior)
        mes_anterior = mnome

    wb.active = 0
    path = "/home/user/Planilha/acompanhamento_assessores.xlsx"
    wb.save(path)
    print(f"Planilha gerada: {path}")


if __name__ == "__main__":
    gerar()
