"""Planilha de Gestão de Carteiras — versão simplificada.

5 abas:
  1. Clientes               — cadastro
  2. Aplicacoes             — base de aplicações
  3. Carteira do Cliente    — visão atual + lista de aplicações
  4. Carteira Sugerida      — editor de realocação (deltas) por Produto
  5. Instruções             — uso

Cada PRODUTO da aba Aplicacoes é uma classe (CDB, LCI, COE, etc.).
"""

import openpyxl
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.worksheet.table import Table, TableStyleInfo

# ─────────────────── CONFIG ───────────────────
ARQUIVO  = "/home/user/Planilha/gestao_carteiras.xlsx"
LIN_CLI  = 1000           # linhas de cadastro
LIN_APL  = 1000           # linhas de aplicações
LIN_HOLD = 60             # itens listados por cliente
CLI_FIM  = LIN_CLI + 2
APL_FIM  = LIN_APL + 2

# Classes = PRODUTOs conhecidos do extrato.
# Editáveis na própria planilha (yellow cells) para acrescentar/remover.
CLASSES = [
    "CDB",
    "LCI",
    "LCA",
    "LIG",
    "POUPANCA",
    "MIN",
    "TESOURO DIRETO",
    "PUBLICO",
    "TESOURARIA  SIGOM",
    "FUNDOS DE INVESTIMENTO",
    "FUNDOS - TORO",
    "ACOES",
    "ETF",
    "FUNDO IMOBILIARIO",
    "TERMO",
    "DIREITO E BONUS",
    "COE",
    "PREVIDENCIA PRIVADA",
]
N_CLASSE = len(CLASSES)

# ─────────────────── ESTILO ───────────────────
NAVY        = '1F3864'
BRANCO      = 'FFFFFF'
CINZA_H     = 'EAEEF3'
AMARELO     = 'FFFBEA'
ZEBRA       = 'F7F9FC'
VERDE       = '1E7B34'
VERDE_LIGHT = 'EAF6EA'
VERM        = 'C0392B'
VERM_LIGHT  = 'FCEBEC'

FMT_BRL = 'R$ #,##0.00;[Red]-R$ #,##0.00'
FMT_PCT = '0.0%;[Red]-0.0%'
FMT_DAT = 'dd/mm/yyyy'

vermelho  = Font(color='9C2B2B', bold=True)
fill_verm = PatternFill('solid', start_color=VERM_LIGHT, end_color=VERM_LIGHT)


def cor(hexcolor):
    return PatternFill('solid', start_color=hexcolor, end_color=hexcolor)


def borda():
    s = Side(style='thin', color='C9D2DD')
    return Border(left=s, right=s, top=s, bottom=s)


def titulo(ws, texto, span):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span)
    c = ws.cell(row=1, column=1, value=texto)
    c.fill = cor(NAVY)
    c.font = Font(bold=True, color=BRANCO, size=15)
    c.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 28


def secao(ws, row, texto, col_ini, col_fim):
    ws.merge_cells(start_row=row, start_column=col_ini,
                   end_row=row, end_column=col_fim)
    c = ws.cell(row=row, column=col_ini, value=texto)
    c.fill = cor(NAVY)
    c.font = Font(bold=True, color=BRANCO, size=11)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[row].height = 22


def cab(ws, row, col, texto):
    c = ws.cell(row=row, column=col, value=texto)
    c.fill = cor(CINZA_H)
    c.border = borda()
    c.font = Font(bold=True, size=10, color=NAVY)
    c.alignment = Alignment(horizontal='center', vertical='center',
                            wrap_text=True)
    return c


def rotulo(ws, row, col, texto):
    c = ws.cell(row=row, column=col, value=texto)
    c.font = Font(bold=True, size=10, color=NAVY)
    c.alignment = Alignment(horizontal='right', vertical='center')
    return c


def campo(ws, row, col, fmt=None, destaque=True):
    c = ws.cell(row=row, column=col)
    c.fill = cor(AMARELO if destaque else BRANCO)
    c.border = borda()
    c.font = Font(size=10)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    if fmt:
        c.number_format = fmt
    return c


# ─────────────────── GERADOR ───────────────────
def construir():
    wb = Workbook()
    wb.remove(wb.active)

    # ─────────────────────────────────────────────────────────────────────
    # ABA 1 — CLIENTES
    # ─────────────────────────────────────────────────────────────────────
    cli = wb.create_sheet("Clientes")
    cli.sheet_view.showGridLines = False
    titulo(cli, "CADASTRO DE CLIENTES", 3)
    cli.column_dimensions['A'].width = 14
    cli.column_dimensions['B'].width = 38
    cli.column_dimensions['C'].width = 18
    cab(cli, 2, 1, "Penumper")
    cab(cli, 2, 2, "Nome")
    cab(cli, 2, 3, "CPF")
    cli.row_dimensions[2].height = 26
    for r in range(3, CLI_FIM + 1):
        zebra = ZEBRA if r % 2 else BRANCO
        for col in (1, 2, 3):
            c = cli.cell(row=r, column=col)
            c.fill = cor(zebra)
            c.border = borda()
            c.font = Font(size=10)
            c.alignment = Alignment(horizontal='left', vertical='center',
                                    indent=1)
            if col in (1, 3):
                c.number_format = '@'
    cli.conditional_formatting.add(
        f"A3:A{CLI_FIM}",
        FormulaRule(formula=[f'AND(A3<>"",COUNTIF($A$3:$A${CLI_FIM},A3)>1)'],
                    font=vermelho, fill=fill_verm))
    cli.conditional_formatting.add(
        f"C3:C{CLI_FIM}",
        FormulaRule(formula=[f'AND(C3<>"",COUNTIF($C$3:$C${CLI_FIM},C3)>1)'],
                    font=vermelho, fill=fill_verm))
    cli.freeze_panes = "A3"

    # ─────────────────────────────────────────────────────────────────────
    # ABA 2 — APLICACOES
    # ─────────────────────────────────────────────────────────────────────
    apl = wb.create_sheet("Aplicacoes")
    apl.sheet_view.showGridLines = False
    titulo(apl, "BASE DE APLICAÇÕES — POSIÇÃO ATUAL", 10)

    colunas = [
        ("Penumper", 14, '@'),
        ("Mês\n(AAAAMM)", 11, '0'),
        ("Produto", 24, None),
        ("Subproduto", 34, None),
        ("Vencimento", 13, FMT_DAT),
        ("Saldo Ponta", 15, FMT_BRL),
        ("Idade", 9, '0'),
        ("Nome\n(automático)", 34, None),
        ("CPF\n(automático)", 18, '@'),
        ("rk", 6, '0'),
    ]
    for i, (txt, larg, _) in enumerate(colunas, 1):
        cab(apl, 2, i, txt)
        apl.column_dimensions[get_column_letter(i)].width = larg
    apl.row_dimensions[2].height = 30

    cli_rng = f"Clientes!$A$3:$C${CLI_FIM}"
    sel_pen = "'Carteira do Cliente'!$B$4"

    for r in range(3, APL_FIM + 1):
        zebra = ZEBRA if r % 2 else BRANCO
        for col in range(1, 11):
            c = apl.cell(row=r, column=col)
            c.fill = cor(zebra)
            c.border = borda()
            c.font = Font(size=10)
            c.alignment = Alignment(horizontal='left', vertical='center',
                                    indent=1)
            fmt = colunas[col - 1][2]
            if fmt:
                c.number_format = fmt
        # Nome (col 8)
        apl.cell(row=r, column=8).value = (
            f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},{cli_rng},2,0),'
            f'"NÃO CADASTRADO"))')
        # CPF (col 9)
        apl.cell(row=r, column=9).value = (
            f'=IF($A{r}="","",IFERROR(VLOOKUP($A{r},{cli_rng},3,0),'
            f'"NÃO CADASTRADO"))')
        # rk (col 10): ranking sequencial das linhas do cliente selecionado
        apl.cell(row=r, column=10).value = (
            f'=IF($A{r}="","",IF($A{r}={sel_pen},'
            f'COUNTIF($A$3:$A{r},{sel_pen}),""))')
        for col in (8, 9):
            apl.cell(row=r, column=col).font = Font(size=10, italic=True,
                                                    color='5B6470')

    apl.conditional_formatting.add(
        f"H3:I{APL_FIM}",
        FormulaRule(formula=['H3="NÃO CADASTRADO"'],
                    font=vermelho, fill=fill_verm))
    apl.column_dimensions['J'].hidden = True
    apl.freeze_panes = "A3"
    tb = Table(displayName="tblAplicacoes", ref=f"A2:J{APL_FIM}")
    tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight9",
                                       showRowStripes=True)
    apl.add_table(tb)

    # Referências utilizadas adiante
    A_PEN = f"Aplicacoes!$A$3:$A${APL_FIM}"
    A_PRD = f"Aplicacoes!$C$3:$C${APL_FIM}"
    A_SUB = f"Aplicacoes!$D$3:$D${APL_FIM}"
    A_VEN = f"Aplicacoes!$E$3:$E${APL_FIM}"
    A_SLD = f"Aplicacoes!$F$3:$F${APL_FIM}"
    A_IDA = f"Aplicacoes!$G$3:$G${APL_FIM}"
    A_RK  = f"Aplicacoes!$J$3:$J${APL_FIM}"

    # ─────────────────────────────────────────────────────────────────────
    # ABA 3 — CARTEIRA DO CLIENTE
    # ─────────────────────────────────────────────────────────────────────
    car = wb.create_sheet("Carteira do Cliente")
    car.sheet_view.showGridLines = False
    titulo(car, "CARTEIRA DO CLIENTE — POSIÇÃO ATUAL", 6)
    for col, larg in ((1, 30), (2, 18), (3, 14),
                      (4, 6), (5, 6), (6, 6)):
        car.column_dimensions[get_column_letter(col)].width = larg

    # Cabeçalho cliente
    rotulo(car, 3, 1, "Cliente:")
    car.merge_cells("B3:D3")
    sel = campo(car, 3, 2)
    sel.font = Font(bold=True, size=11, color=NAVY)
    rotulo(car, 4, 1, "Penumper:")
    rotulo(car, 5, 1, "CPF:")
    rotulo(car, 6, 1, "Patrimônio total:")
    rotulo(car, 7, 1, "Nº de aplicações:")

    cli_nome = f"Clientes!$B$3:$B${CLI_FIM}"
    cli_pen  = f"Clientes!$A$3:$A${CLI_FIM}"
    cli_cpf  = f"Clientes!$C$3:$C${CLI_FIM}"
    for r, formula, fmt in (
            (4, f'=IFERROR(INDEX({cli_pen},MATCH($B$3,{cli_nome},0)),"")', '@'),
            (5, f'=IFERROR(INDEX({cli_cpf},MATCH($B$3,{cli_nome},0)),"")', '@'),
            (6, f'=SUMIFS({A_SLD},{A_PEN},$B$4)', FMT_BRL),
            (7, f'=COUNTIF({A_PEN},$B$4)', '0')):
        c = campo(car, r, 2, fmt, destaque=False)
        c.value = formula
        c.font = Font(bold=True, size=10, color=NAVY)

    dv = DataValidation(type="list", formula1=f"={cli_nome}",
                        allow_blank=True,
                        prompt="Escolha o cliente na lista")
    car.add_data_validation(dv)
    dv.add(car["B3"])

    # ---- alocação por produto ------------------------------------------
    LIN_CL0 = 11
    secao(car, 9, "ALOCAÇÃO POR PRODUTO (atual)", 1, 6)
    for col, txt in ((1, "Produto"), (2, "Saldo (R$)"), (3, "% atual")):
        cab(car, 10, col, txt)
    for i, lab in enumerate(CLASSES):
        r = LIN_CL0 + i
        zebra = ZEBRA if i % 2 else BRANCO
        ca = car.cell(row=r, column=1, value=lab)
        ca.fill = cor(AMARELO)          # editável (o usuário pode trocar/acrescentar produtos)
        ca.border = borda()
        ca.font = Font(size=10, bold=True, color=NAVY)
        ca.alignment = Alignment(horizontal='left', vertical='center',
                                 indent=1)
        cb = car.cell(row=r, column=2,
                      value=f'=SUMIFS({A_SLD},{A_PEN},$B$4,{A_PRD},$A{r})')
        cb.fill = cor(zebra); cb.border = borda(); cb.font = Font(size=10)
        cb.alignment = Alignment(horizontal='right', vertical='center')
        cb.number_format = FMT_BRL
        cc = car.cell(row=r, column=3,
                      value=f'=IFERROR($B{r}/$B$6,0)')
        cc.fill = cor(zebra); cc.border = borda(); cc.font = Font(size=10)
        cc.alignment = Alignment(horizontal='right', vertical='center')
        cc.number_format = FMT_PCT
    r_tot = LIN_CL0 + N_CLASSE
    r_out = r_tot + 1
    for r, lab, f2, f3 in (
        (r_tot, "TOTAL CLASSIFICADO",
         f'=SUM(B{LIN_CL0}:B{r_tot-1})',
         f'=SUM(C{LIN_CL0}:C{r_tot-1})'),
        (r_out, "Outros (produtos fora da lista)",
         f'=$B$6-$B${r_tot}', f'=IFERROR($B{r_out}/$B$6,0)')):
        ca = car.cell(row=r, column=1, value=lab)
        cb = car.cell(row=r, column=2, value=f2)
        cb.number_format = FMT_BRL
        cc = car.cell(row=r, column=3, value=f3)
        cc.number_format = FMT_PCT
        for col in (1, 2, 3):
            ccc = car.cell(row=r, column=col)
            ccc.fill = cor(CINZA_H)
            ccc.border = borda()
            ccc.font = Font(bold=True, size=10, color=NAVY)
            ccc.alignment = Alignment(
                horizontal='left' if col == 1 else 'right',
                vertical='center', indent=1 if col == 1 else 0)

    # destaca a linha "Outros" se houver saldo lá
    car.conditional_formatting.add(
        f"A{r_out}:C{r_out}",
        FormulaRule(formula=[f'$B${r_out}>0.5'],
                    font=Font(bold=True, color='B0301A'),
                    fill=fill_verm))

    # gráfico de rosca
    rosca = DoughnutChart()
    rosca.title = "Alocação por produto"
    rosca.height, rosca.width = 10, 13
    dados = Reference(car, min_col=2, min_row=10, max_row=r_tot - 1)
    cats  = Reference(car, min_col=1, min_row=LIN_CL0, max_row=r_tot - 1)
    rosca.add_data(dados, titles_from_data=True)
    rosca.set_categories(cats)
    rosca.dataLabels = DataLabelList()
    rosca.dataLabels.showPercent = True
    car.add_chart(rosca, "F3")

    # ---- lista de aplicações do cliente --------------------------------
    LIN_H0 = r_out + 3
    secao(car, LIN_H0 - 1, "APLICAÇÕES DO CLIENTE", 1, 6)
    for col, txt in ((1, "Produto"), (2, "Subproduto"), (3, "Vencimento"),
                     (4, "Saldo (R$)"), (5, "% carteira"), (6, "Idade")):
        cab(car, LIN_H0, col, txt)
    for k in range(1, LIN_HOLD + 1):
        r = LIN_H0 + k
        zebra = ZEBRA if k % 2 else BRANCO
        defs = [
            (1, f'=IFERROR(INDEX({A_PRD},MATCH({k},{A_RK},0)),"")', None),
            (2, f'=IFERROR(INDEX({A_SUB},MATCH({k},{A_RK},0)),"")', None),
            (3, f'=IFERROR(INDEX({A_VEN},MATCH({k},{A_RK},0)),"")', FMT_DAT),
            (4, f'=IFERROR(INDEX({A_SLD},MATCH({k},{A_RK},0)),"")', FMT_BRL),
            (5, f'=IFERROR(INDEX({A_SLD},MATCH({k},{A_RK},0))/$B$6,"")',
             FMT_PCT),
            (6, f'=IFERROR(INDEX({A_IDA},MATCH({k},{A_RK},0)),"")', '0'),
        ]
        for col, formula, fmt in defs:
            c = car.cell(row=r, column=col, value=formula)
            c.fill = cor(zebra)
            c.border = borda()
            c.font = Font(size=10)
            align = 'right' if col in (3, 4, 5, 6) else 'left'
            c.alignment = Alignment(horizontal=align, vertical='center',
                                    indent=1 if align == 'left' else 0)
            if fmt:
                c.number_format = fmt

    car.freeze_panes = "A2"

    # ─────────────────────────────────────────────────────────────────────
    # ABA 4 — CARTEIRA SUGERIDA (editor de realocação por delta)
    # ─────────────────────────────────────────────────────────────────────
    sug = wb.create_sheet("Carteira Sugerida")
    sug.sheet_view.showGridLines = False
    titulo(sug, "REALOCAÇÃO DA CARTEIRA", 14)
    larguras_sug = [30, 18, 18, 18, 12, 4,
                    14, 14, 14, 14, 14, 14, 14, 14]
    for i, larg in enumerate(larguras_sug, 1):
        sug.column_dimensions[get_column_letter(i)].width = larg

    # cabeçalho cliente
    rotulo(sug, 3, 1, "Cliente:")
    sug.merge_cells("B3:E3")
    c = campo(sug, 3, 2, destaque=False)
    c.value = "='Carteira do Cliente'!B3"
    c.font = Font(bold=True, size=12, color=NAVY)
    rotulo(sug, 4, 1, "Patrimônio Total:")
    c = campo(sug, 4, 2, FMT_BRL, destaque=False)
    c.value = "='Carteira do Cliente'!B6"
    c.font = Font(bold=True, size=10, color=NAVY)

    # banner: soma dos deltas — deve fechar em zero
    DATA_FIM = 7 + N_CLASSE          # última linha de dados (8..7+N)
    sug.merge_cells("G3:N3")
    c = sug.cell(row=3, column=7, value="REALOCAÇÃO LÍQUIDA (DEVE SER R$ 0,00)")
    c.fill = cor(NAVY); c.font = Font(bold=True, color=BRANCO, size=11)
    c.alignment = Alignment(horizontal='center', vertical='center')
    sug.merge_cells("G4:N4")
    bal = sug.cell(row=4, column=7, value=f"=SUM(C8:C{DATA_FIM})")
    bal.fill = cor(VERDE)
    bal.font = Font(bold=True, color=BRANCO, size=18)
    bal.alignment = Alignment(horizontal='center', vertical='center')
    bal.number_format = FMT_BRL
    sug.row_dimensions[3].height = 22
    sug.row_dimensions[4].height = 32

    # conditional: vermelho se desbalanceado
    sug.conditional_formatting.add(
        "G4:N4",
        FormulaRule(formula=[f'ROUND(SUM($C$8:$C${DATA_FIM}),2)<>0'],
                    fill=cor(VERM)))

    # tabela de realocação
    secao(sug, 6, "REALOCAÇÃO POR PRODUTO", 1, 5)
    heads_sug = ["Produto", "Saldo Atual (R$)",
                 "Delta (R$)\n+ aplicar / - reduzir",
                 "Saldo Novo (R$)", "% Novo"]
    for col, txt in enumerate(heads_sug, 1):
        cab(sug, 7, col, txt)
    sug.row_dimensions[7].height = 32

    for i, _ in enumerate(CLASSES):
        r = 8 + i                  # 8 .. 8+N-1
        zebra = ZEBRA if i % 2 else BRANCO
        # A: Produto (mirror da Carteira do Cliente)
        ca = sug.cell(row=r, column=1,
                      value=f"='Carteira do Cliente'!A{LIN_CL0+i}")
        ca.fill = cor(zebra); ca.border = borda()
        ca.font = Font(size=10, bold=True, color=NAVY)
        ca.alignment = Alignment(horizontal='left', vertical='center',
                                 indent=1)
        # B: Saldo Atual (mirror)
        cb = sug.cell(row=r, column=2,
                      value=f"='Carteira do Cliente'!B{LIN_CL0+i}")
        cb.fill = cor(zebra); cb.border = borda(); cb.font = Font(size=10)
        cb.alignment = Alignment(horizontal='right', vertical='center')
        cb.number_format = FMT_BRL
        # C: Delta (input)
        cc = sug.cell(row=r, column=3)
        cc.fill = cor(AMARELO); cc.border = borda()
        cc.font = Font(size=11, bold=True)
        cc.alignment = Alignment(horizontal='right', vertical='center')
        cc.number_format = FMT_BRL
        # D: Saldo Novo
        cd = sug.cell(row=r, column=4, value=f"=B{r}+C{r}")
        cd.fill = cor(zebra); cd.border = borda()
        cd.font = Font(size=10, bold=True, color=NAVY)
        cd.alignment = Alignment(horizontal='right', vertical='center')
        cd.number_format = FMT_BRL
        # E: % Novo
        ce = sug.cell(row=r, column=5,
                      value=f"=IFERROR($D{r}/SUM($D$8:$D${DATA_FIM}),0)")
        ce.fill = cor(zebra); ce.border = borda(); ce.font = Font(size=10)
        ce.alignment = Alignment(horizontal='right', vertical='center')
        ce.number_format = FMT_PCT

    # linha TOTAL
    r_sug_tot = DATA_FIM + 1        # 8+N
    for col, val in ((1, "TOTAL"),
                     (2, f"=SUM(B8:B{DATA_FIM})"),
                     (3, f"=SUM(C8:C{DATA_FIM})"),
                     (4, f"=SUM(D8:D{DATA_FIM})"),
                     (5, f"=SUM(E8:E{DATA_FIM})")):
        c = sug.cell(row=r_sug_tot, column=col, value=val)
        c.fill = cor(CINZA_H); c.border = borda()
        c.font = Font(bold=True, size=10, color=NAVY)
        c.alignment = Alignment(
            horizontal='left' if col == 1 else 'right',
            vertical='center', indent=1 if col == 1 else 0)
        if col in (2, 3, 4):
            c.number_format = FMT_BRL
        elif col == 5:
            c.number_format = FMT_PCT

    # destaca deltas: verde positivos, vermelho negativos
    sug.conditional_formatting.add(
        f"C8:C{DATA_FIM}",
        FormulaRule(formula=[f'C8<0'],
                    font=Font(color=VERM, bold=True)))
    sug.conditional_formatting.add(
        f"C8:C{DATA_FIM}",
        FormulaRule(formula=[f'C8>0'],
                    font=Font(color=VERDE, bold=True)))
    # destaca saldo novo negativo (não pode subtrair mais do que tem)
    sug.conditional_formatting.add(
        f"D8:D{DATA_FIM}",
        FormulaRule(formula=[f'D8<0'],
                    font=vermelho, fill=fill_verm))

    # gráficos: rosca Atual e Proposta lado a lado
    rosca_at = DoughnutChart()
    rosca_at.title = "Alocação ATUAL"
    rosca_at.height, rosca_at.width = 9, 11
    rosca_at.add_data(Reference(sug, min_col=2, min_row=7,
                                 max_row=DATA_FIM),
                      titles_from_data=True)
    rosca_at.set_categories(Reference(sug, min_col=1, min_row=8,
                                       max_row=DATA_FIM))
    rosca_at.dataLabels = DataLabelList()
    rosca_at.dataLabels.showPercent = True
    sug.add_chart(rosca_at, "G6")

    rosca_pr = DoughnutChart()
    rosca_pr.title = "Alocação PROPOSTA"
    rosca_pr.height, rosca_pr.width = 9, 11
    rosca_pr.add_data(Reference(sug, min_col=4, min_row=7,
                                 max_row=DATA_FIM),
                      titles_from_data=True)
    rosca_pr.set_categories(Reference(sug, min_col=1, min_row=8,
                                       max_row=DATA_FIM))
    rosca_pr.dataLabels = DataLabelList()
    rosca_pr.dataLabels.showPercent = True
    sug.add_chart(rosca_pr, "G24")

    # observações livres
    R_OBS = r_sug_tot + 3
    secao(sug, R_OBS, "OBSERVAÇÕES / RACIONAL DA REALOCAÇÃO", 1, 14)
    for r in range(R_OBS + 1, R_OBS + 5):
        sug.merge_cells(start_row=r, start_column=1,
                        end_row=r, end_column=14)
        c = sug.cell(row=r, column=1)
        c.fill = cor(AMARELO); c.border = borda()
        c.font = Font(size=10)
        c.alignment = Alignment(horizontal='left', vertical='top',
                                wrap_text=True, indent=1)
        sug.row_dimensions[r].height = 24

    sug.freeze_panes = "A8"

    # ─────────────────────────────────────────────────────────────────────
    # ABA 5 — INSTRUÇÕES
    # ─────────────────────────────────────────────────────────────────────
    ins = wb.create_sheet("Instruções")
    ins.sheet_view.showGridLines = False
    titulo(ins, "INSTRUÇÕES DE USO", 6)
    ins.column_dimensions['A'].width = 110

    linhas = [
        "",
        "1. Aba 'Clientes' — cadastre Penumper, Nome e CPF de cada cliente. "
        "Duplicatas (Penumper ou CPF) ficam destacadas em vermelho.",
        "",
        "2. Aba 'Aplicacoes' — cole a posição vinda do sistema: Penumper, "
        "Mês, Produto, Subproduto, Vencimento, Saldo Ponta e Idade. As "
        "colunas Nome e CPF se preenchem sozinhas via Penumper. "
        "'NÃO CADASTRADO' significa que o Penumper não está na aba Clientes.",
        "",
        "3. Aba 'Carteira do Cliente' — escolha o cliente no dropdown. "
        "A planilha mostra o patrimônio total, a alocação por Produto (CDB, "
        "LCI, COE, etc.) e a lista de aplicações daquele cliente. "
        "Se a linha 'Outros' acender em vermelho, é porque há saldo num "
        "Produto que não está na lista — acrescente a linha do produto na "
        "tabela de alocação (a coluna A é editável).",
        "",
        "4. Aba 'Carteira Sugerida' — editor de realocação. Para cada "
        "Produto digite na coluna 'Delta' o quanto quer aumentar (+) ou "
        "diminuir (-). 'Saldo Novo' e '% Novo' se atualizam sozinhos, "
        "assim como os gráficos 'Atual' e 'Proposta'.",
        "     O banner 'Realocação Líquida' fica VERDE quando a soma dos "
        "deltas dá zero (você tirou de um lado e colocou em outro do "
        "mesmo tamanho) e VERMELHO enquanto não fechar.",
        "     Use o bloco 'Observações / Racional' para anotar o motivo "
        "das movimentações (ex.: tirei R$ 20k de CDB porque vence; coloquei "
        "em LCA isenta de 11,5% a.a.).",
        "",
        "Observação: a planilha trabalha com a posição atual (foto). "
        "Para refletir uma nova posição, basta colar mais linhas na aba "
        "'Aplicacoes' — todas as fórmulas se ajustam.",
    ]
    for i, txt in enumerate(linhas, 2):
        c = ins.cell(row=i, column=1, value=txt)
        c.alignment = Alignment(horizontal='left', vertical='center',
                                wrap_text=True)
        neg = txt[:2].strip().rstrip('.').isdigit()
        c.font = Font(size=11, bold=neg, color=NAVY if neg else '333333')
        ins.row_dimensions[i].height = 22

    wb.active = car
    wb.save(ARQUIVO)
    print(f"Planilha gerada: {ARQUIVO}")


if __name__ == "__main__":
    construir()
